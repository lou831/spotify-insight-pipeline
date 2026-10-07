"""Build golden_50_labeling.xlsx: the 50 golden reviews with dropdowns for hand labels.
Stars are hidden on purpose so they do not sway your labels (the model never sees them either).

    python pipeline/make_label_sheet.py
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import INTENTS, ROOT, TOPICS, csv_rows

SRC = ROOT / "golden_50_to_label.csv"
DST = ROOT / "golden_50_labeling.xlsx"
COLS = ["#", "review_id", "review_text", "topic", "intent", "sentiment", "severity",
        "entities", "evidence_quote", "needs_review", "my_notes"]
WIDTHS = [4, 12, 70, 13, 14, 11, 10, 20, 40, 14, 40]


def main():
    if DST.exists():
        raise SystemExit(f"{DST.name} already exists; not overwriting your labels.")
    wb = Workbook()
    ws = wb.active
    ws.title = "label_here"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for i, r in enumerate(csv_rows(SRC), 1):
        ws.append([i, r["review_id"], r["review_text"]] + [""] * 8)
    n = ws.max_row
    for col, width in zip("ABCDEFGHIJK", WIDTHS):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"

    def dropdown(col, values, title):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle=title, error="Pick from the list.")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")

    dropdown("D", TOPICS, "topic")
    dropdown("E", INTENTS, "intent")
    dropdown("G", ["1", "2", "3", "4", "5"], "severity")
    dropdown("J", ["FALSE", "TRUE"], "needs_review")
    sv = DataValidation(type="decimal", operator="between", formula1="-1", formula2="1",
                        showErrorMessage=True, error="Enter a number from -1 to 1.")
    ws.add_data_validation(sv)
    sv.add(f"F2:F{n}")

    ref = wb.create_sheet("cheat_sheet")
    lines = [
        "TOPIC: access=login/account | usability=navigation, controls, layout, queue, ADS | playback=won't play, crash, lag, audio, battery",
        "       downloads=offline/saved music | catalog=missing songs, search, recommendations, lyrics",
        "       billing=price, charges, subscription, paywall, 'need premium to skip' | support=customer service | other=general praise/criticism",
        "INTENT (first that fits): cancellation > complaint > request > praise > unclear",
        "SEVERITY: 1 no problem/praise/request | 2 annoyance, generic 'bad app' | 3 feature degraded, workaround exists",
        "          4 core task blocked (can't log in / play / open) | 5 explicit serious money, privacy or data harm ONLY",
        "SENTIMENT: -1 very negative ... 0 neutral ... +1 very positive (one decimal is fine)",
        "ENTITIES: features actually named, comma-separated (e.g. shuffle, ads). Leave blank if none.",
        "EVIDENCE_QUOTE: copy-paste exact words from the review (select inside the cell, Cmd+C). Short is fine.",
        "NEEDS_REVIEW: TRUE if ambiguous, missing context, or a language you can't read.",
        "Tie-breaks: highest severity problem wins; then the first one mentioned. Premium mention alone != billing.",
    ]
    for line in lines:
        ref.append([line])
    ref.column_dimensions["A"].width = 150
    wb.save(DST)
    print(f"wrote {DST.name} with {n - 1} reviews")


if __name__ == "__main__":
    main()
