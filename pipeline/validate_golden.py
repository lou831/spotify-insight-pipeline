"""Check your hand labels and export them to evals/golden_50_labeled.csv (no model calls).

    python pipeline/validate_golden.py

Reports blanks, invalid values, and evidence quotes that are not exact substrings.
Run it as often as you like while labeling.
"""

import csv
import json

from openpyxl import load_workbook

from common import INTENTS, ROOT, TOPICS, csv_rows

SHEET = ROOT / "golden_50_labeling.xlsx"
OUTFILE = ROOT / "evals" / "golden_50_labeled.csv"
FIELDS = ["topic", "intent", "sentiment", "severity", "entities", "evidence_quote", "needs_review"]


def cell(v):
    return "" if v is None else str(v).strip()


def main():
    source = {r["review_id"]: r for r in csv_rows(ROOT / "golden_50_to_label.csv")}
    ws = load_workbook(SHEET, data_only=True)["label_here"]
    header = [cell(c.value) for c in ws[1]]
    problems, out, complete = [], [], 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(header, row))
        rid = cell(r["review_id"])
        if not rid:
            continue
        num = r["#"]
        text = source[rid]["review_text"]
        errs = []
        topic, intent = cell(r["topic"]), cell(r["intent"])
        if topic not in TOPICS:
            errs.append("topic blank/invalid")
        if intent not in INTENTS:
            errs.append("intent blank/invalid")
        try:
            sent = float(cell(r["sentiment"]))
            if not -1 <= sent <= 1:
                errs.append("sentiment outside -1..1")
        except ValueError:
            sent = None
            errs.append("sentiment blank/not a number")
        sev_s = cell(r["severity"]).removesuffix(".0")
        sev = int(sev_s) if sev_s in {"1", "2", "3", "4", "5"} else None
        if sev is None:
            errs.append("severity blank/invalid")
        quote = "" if r["evidence_quote"] is None else str(r["evidence_quote"])
        if not quote.strip():
            errs.append("evidence_quote blank")
        elif quote not in text:
            stripped = quote.strip()
            if stripped and stripped in text:
                quote = stripped  # trailing spaces from copy-paste
            else:
                errs.append("evidence_quote is not an exact substring of the review")
        nr = cell(r["needs_review"]).upper()
        if nr not in {"TRUE", "FALSE"}:
            errs.append("needs_review blank")
        if intent in {"praise", "request", "unclear"} and sev and sev > 1:
            errs.append(f"note: intent={intent} usually means severity 1 (you put {sev})")
        if intent == "complaint" and sev == 1:
            errs.append("note: a complaint is usually severity 2+")
        ents = [e.strip() for e in cell(r["entities"]).split(",") if e.strip()]
        if errs:
            problems.append(f"#{num}: " + "; ".join(errs))
        if not [e for e in errs if not e.startswith("note:")]:
            complete += 1
        out.append({"review_id": rid, "topic": topic, "intent": intent, "sentiment": "" if sent is None else sent,
                    "severity": "" if sev is None else sev, "entities": json.dumps(ents, ensure_ascii=False),
                    "evidence_quote": quote, "needs_review": nr, "notes": cell(r.get("my_notes"))})
    OUTFILE.parent.mkdir(exist_ok=True)
    with OUTFILE.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["review_id"] + FIELDS + ["notes"])
        w.writeheader()
        w.writerows(out)
    print(f"{complete}/{len(out)} reviews fully labeled. Saved {OUTFILE.relative_to(ROOT)}")
    for p in problems:
        print("  " + p)


if __name__ == "__main__":
    main()
